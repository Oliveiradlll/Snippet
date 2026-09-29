#Disparador#

import json
import os
import re
import sys
import time
import urllib.parse
import webbrowser
from datetime import datetime, timedelta
import customtkinter as ctk
from tkinter import messagebox, simpledialog, filedialog

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    import keyboard
    KEYBOARD_AVAILABLE = True
except ImportError:
    KEYBOARD_AVAILABLE = False

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

FONT_MAIN = ("Segoe UI Variable Display", 11)
FONT_BOLD = ("Segoe UI Variable Display", 11, "bold")
FONT_TITLE = ("Segoe UI Variable Display", 14, "bold")
FONT_SMALL = ("Segoe UI Variable Text", 10)
FONT_BADGE = ("Segoe UI Variable Text", 9, "bold")

DEFAULT_DATA = {
  "schemaVersion": 1,
  "settings": {
      "theme": "dark"
  },
  "categories": [
    {
      "id": "campanhas",
      "name": "Campanhas de Disparo",
      "snippets": [
        {"id": "oferta-especial", "label": "Oferta Especial", "shortcut": "CTRL+1", "body": "[SAUDACAO]! Preparamos uma condição imperdível para o seu atendimento hoje, [HOJE]."},
        {"id": "pesquisa-satisfacao", "label": "Pesquisa de Satisfação", "shortcut": "CTRL+2", "body": "[SAUDACAO]! Como foi a sua experiência conosco? Sua opinião é muito importante."},
        {"id": "lembrete-geral", "label": "Lembrete de Retorno", "shortcut": "CTRL+3", "body": "[SAUDACAO]! Passando para saber se conseguiu avaliar nossa proposta anterior."}
      ]
    }
  ]
}

def position_side_window(parent, dialog, dialog_width=380, dialog_height=500):
    parent_x = parent.winfo_x()
    parent_y = parent.winfo_y()
    parent_width = parent.winfo_width()
    
    target_x = parent_x + parent_width + 12
    target_y = parent_y

    dialog.geometry(f"{dialog_width}x{dialog_height}+{target_x}+{target_y}")

def get_data_filepath():
    if getattr(sys, 'frozen', False):
        application_path = os.path.dirname(sys.executable)
    else:
        application_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(application_path, "snippets_mass.json")

def get_history_filepath():
    if getattr(sys, 'frozen', False):
        application_path = os.path.dirname(sys.executable)
    else:
        application_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(application_path, "history_mass.json")

def load_data():
    filepath = get_data_filepath()
    if os.path.exists(filepath):
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                data = json.load(f)
                if "settings" not in data:
                    data["settings"] = DEFAULT_DATA["settings"]
                return data
        except Exception:
            return DEFAULT_DATA
    return DEFAULT_DATA

def save_data(data):
    filepath = get_data_filepath()
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def load_history():
    filepath = get_history_filepath()
    if os.path.exists(filepath):
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_history(history_data):
    filepath = get_history_filepath()
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(history_data, f, ensure_ascii=False, indent=2)

class HistoryDialog(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("📊 Relatório de Disparos")
        self.attributes('-topmost', True)
        self.parent_app = parent
        parent.active_history_dialog = self

        self.create_widgets()
        position_side_window(parent, self, 380, 500)

    def create_widgets(self):
        for widget in self.winfo_children():
            widget.destroy()

        today_date = datetime.now().strftime("%Y-%m-%d")
        display_date = datetime.now().strftime("%d/%m/%Y")
        
        history = load_history()
        day_data = history.get(today_date, {"total": 0, "snippets": {}, "numbers": []})

        lbl_title = ctk.CTkLabel(self, text=f"Relatório de Hoje ({display_date})", font=FONT_TITLE)
        lbl_title.pack(pady=(20, 10))

        self.lbl_total = ctk.CTkLabel(self, text=f"Total de Disparos: {day_data['total']}", font=("Segoe UI Variable Display", 16, "bold"), text_color="#2ecc71")
        self.lbl_total.pack(pady=(0, 15))

        ctk.CTkLabel(self, text="Mensagens Utilizadas:", font=FONT_BOLD).pack(anchor="w", padx=20)
        self.frame_stats = ctk.CTkScrollableFrame(self, height=120)
        self.frame_stats.pack(fill="x", padx=20, pady=(5, 15))

        if not day_data["snippets"]:
            ctk.CTkLabel(self.frame_stats, text="Nenhum disparo registrado hoje.").pack(pady=10)
        else:
            for label, count in day_data["snippets"].items():
                row = ctk.CTkFrame(self.frame_stats, fg_color="transparent")
                row.pack(fill="x", pady=2)
                ctk.CTkLabel(row, text=label, font=FONT_MAIN).pack(side="left")
                ctk.CTkLabel(row, text=str(count), font=FONT_BOLD).pack(side="right")

        ctk.CTkLabel(self, text="Contatos Atingidos:", font=FONT_BOLD).pack(anchor="w", padx=20)
        
        self.txt_nums = ctk.CTkTextbox(self, height=150, font=FONT_MAIN, corner_radius=8)
        self.txt_nums.pack(fill="both", expand=True, padx=20, pady=(5, 20))

        if day_data["numbers"]:
            self.txt_nums.insert("1.0", "\n".join(day_data["numbers"]))
        else:
            self.txt_nums.insert("1.0", "Nenhum número registrado hoje.")
        
        self.txt_nums.configure(state="disabled")

    def refresh_data(self):
        try:
            if self.winfo_exists():
                self.create_widgets()
        except Exception:
            pass

    def destroy(self):
        if hasattr(self.parent_app, 'active_history_dialog'):
            self.parent_app.active_history_dialog = None
        super().destroy()

class FillSnippetDialog(ctk.CTkToplevel):
    def __init__(self, parent, label, body_template):
        super().__init__(parent)
        self.title(f"Variáveis: {label}")
        self.attributes('-topmost', True)
        self.grab_set()
        self.result_text = None

        self.body_template = body_template
        self.placeholders = list(dict.fromkeys(re.findall(r'\[(.*?)\]', body_template)))

        lbl = ctk.CTkLabel(self, text="Preencha os campos dinâmicos:", font=FONT_TITLE)
        lbl.pack(anchor="w", padx=20, pady=(20, 10))

        self.entries = {}
        fields_frame = ctk.CTkScrollableFrame(self)
        fields_frame.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        first_entry = None
        for ph in self.placeholders:
            frame_item = ctk.CTkFrame(fields_frame, fg_color="transparent")
            frame_item.pack(fill="x", pady=6)
            lbl_field = ctk.CTkLabel(frame_item, text=f"[{ph}]", font=FONT_BOLD, width=90, anchor="w")
            lbl_field.pack(side="left")
            ent = ctk.CTkEntry(frame_item, corner_radius=8, font=FONT_MAIN)
            ent.pack(side="right", fill="x", expand=True)
            ent.bind("<Return>", lambda e: self.on_confirm())
            self.entries[ph] = ent
            if first_entry is None: first_entry = ent

        btn_copy = ctk.CTkButton(self, text="🚀 Confirmar e Disparar", font=FONT_BOLD, fg_color="#2ecc71", hover_color="#27ae60", corner_radius=8, height=40, command=self.on_confirm)
        btn_copy.pack(fill="x", padx=20, pady=16)
        position_side_window(parent, self, 360, 420)
        if first_entry: first_entry.focus_set()

    def on_confirm(self):
        final_text = self.body_template
        for ph, entry in self.entries.items():
            val = entry.get().strip()
            final_text = final_text.replace(f"[{ph}]", val if val else f"[{ph}]")
        self.result_text = final_text
        self.destroy()

class EditSnippetDialog(ctk.CTkToplevel):
    def __init__(self, parent, label="", body="", shortcut=""):
        super().__init__(parent)
        self.title("Gerenciar Mensagem")
        self.attributes('-topmost', True)
        self.grab_set()
        self.saved = False

        main_frame = ctk.CTkFrame(self, fg_color="transparent")
        main_frame.pack(fill="both", expand=True, padx=16, pady=16)

        ctk.CTkLabel(main_frame, text="Título da Mensagem", font=FONT_BOLD).pack(anchor="w", pady=(0, 4))
        self.ent_label = ctk.CTkEntry(main_frame, corner_radius=8, font=FONT_MAIN)
        self.ent_label.insert(0, label)
        self.ent_label.pack(fill="x", pady=(0, 12))

        ctk.CTkLabel(main_frame, text="Atalho Opcional (ex: CTRL+5)", font=FONT_BOLD).pack(anchor="w", pady=(0, 4))
        self.ent_shortcut = ctk.CTkEntry(main_frame, corner_radius=8, font=FONT_MAIN)
        self.ent_shortcut.insert(0, shortcut)
        self.ent_shortcut.pack(fill="x", pady=(0, 12))

        ctk.CTkLabel(main_frame, text="Texto do Disparo ([SAUDACAO], [HOJE], [VAR])", font=FONT_BOLD).pack(anchor="w", pady=(0, 4))
        self.txt_body = ctk.CTkTextbox(main_frame, corner_radius=8, font=FONT_MAIN, height=120)
        self.txt_body.insert("1.0", body)
        self.txt_body.pack(fill="both", expand=True, pady=(0, 16))

        btn_save = ctk.CTkButton(main_frame, text="💾 Salvar Mensagem", font=FONT_BOLD, corner_radius=8, height=40, command=self.on_save)
        btn_save.pack(fill="x")
        position_side_window(parent, self, 360, 440)

    def on_save(self):
        self.label_val = self.ent_label.get().strip()
        self.shortcut_val = self.ent_shortcut.get().strip().lower()
        self.body_val = self.txt_body.get("1.0", "end-1c").strip()
        if not self.label_val or not self.body_val:
            messagebox.showwarning("Aviso", "Preencha o título e o texto.", parent=self)
            return
        self.saved = True
        self.destroy()

class MassSenderApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.data = load_data()
        self.current_theme = self.data["settings"].get("theme", "dark")
        self.active_history_dialog = None
        self.is_visible = True

        self.title("ZapBatch - Disparador em Massa")
        self.geometry("480x790")
        self.attributes('-topmost', True)

        header_frame = ctk.CTkFrame(self, fg_color="transparent")
        header_frame.pack(fill="x", padx=16, pady=(14, 6))

        top_bar = ctk.CTkFrame(header_frame, fg_color="transparent")
        top_bar.pack(fill="x", pady=(0, 6))

        ctk.CTkLabel(top_bar, text="🚀 ZapBatch ⚡", font=FONT_TITLE, text_color="#f1c40f").pack(side="left")
        
        ctk.CTkButton(top_bar, text="📊 Relatório", font=FONT_BOLD, width=85, height=30, corner_radius=8, command=self.show_history).pack(side="right")
        ctk.CTkButton(top_bar, text="◐", font=FONT_BOLD, width=35, height=30, corner_radius=8, command=self.toggle_theme).pack(side="right", padx=(0, 6))

        box_nums = ctk.CTkFrame(self, corner_radius=10)
        box_nums.pack(fill="x", padx=16, pady=(2, 8))

        top_col_header = ctk.CTkFrame(box_nums, fg_color="transparent")
        top_col_header.pack(fill="x", padx=12, pady=(8, 2))
        
        ctk.CTkLabel(top_col_header, text="📋 1. Destinatários (Cole ou Importe)", font=FONT_BOLD).pack(side="left")
        
        actions_right = ctk.CTkFrame(top_col_header, fg_color="transparent")
        actions_right.pack(side="right")

        ctk.CTkButton(actions_right, text="🗑️ Limpar", font=FONT_BOLD, width=70, height=26, fg_color="#c0392b", hover_color="#a93226", corner_radius=6, command=self.clear_recipients).pack(side="left", padx=(0, 6))
        ctk.CTkButton(actions_right, text="📁 Importar", font=FONT_BOLD, width=90, height=26, fg_color="#27ae60", hover_color="#219653", corner_radius=6, command=self.import_spreadsheet).pack(side="left")

        self.txt_whatsapp_batch = ctk.CTkTextbox(box_nums, height=70, font=FONT_MAIN, corner_radius=8)
        self.txt_whatsapp_batch.pack(fill="x", padx=12, pady=(0, 4))
        self.txt_whatsapp_batch.insert("1.0", "")
        self.txt_whatsapp_batch.bind("<KeyRelease>", self.update_parsed_numbers)

        ctk.CTkButton(box_nums, text="➕ Importar Número para a Lista", font=FONT_BOLD, height=28, corner_radius=6, fg_color="#2980b9", hover_color="#2471a3", command=self.add_number_to_list).pack(fill="x", padx=12, pady=(0, 6))

        table_frame = ctk.CTkFrame(box_nums, fg_color=("gray90", "gray17"), corner_radius=8)
        table_frame.pack(fill="x", padx=12, pady=(0, 10))

        sub_table_header = ctk.CTkFrame(table_frame, fg_color="transparent")
        sub_table_header.pack(fill="x", padx=8, pady=(4, 0))
        ctk.CTkLabel(sub_table_header, text="🔍 Números Carregados:", font=FONT_SMALL, text_color=("gray40", "gray60")).pack(side="left")
        self.lbl_count_valid = ctk.CTkLabel(sub_table_header, text="0 números", font=FONT_BADGE, text_color="#2ecc71")
        self.lbl_count_valid.pack(side="right")
        
        self.scroll_column_view = ctk.CTkScrollableFrame(table_frame, height=65, fg_color="transparent")
        self.scroll_column_view.pack(fill="x", padx=4, pady=(2, 6))

        box_msgs = ctk.CTkFrame(self, fg_color="transparent")
        box_msgs.pack(fill="both", expand=True, padx=16, pady=(0, 8))

        cat_row = ctk.CTkFrame(box_msgs, fg_color="transparent")
        cat_row.pack(fill="x", pady=(0, 4))

        self.combo_cat = ctk.CTkOptionMenu(cat_row, font=FONT_MAIN, dynamic_resizing=False, command=self.on_category_change, corner_radius=8)
        self.combo_cat.pack(side="left", fill="x", expand=True)

        ctk.CTkButton(cat_row, text="+ Grupo", width=70, height=32, font=FONT_BOLD, fg_color="#2ecc71", hover_color="#27ae60", corner_radius=8, command=self.add_category).pack(side="left", padx=(6, 0))

        ctk.CTkLabel(box_msgs, text="💬 2. Escolha o Modelo de Mensagem para Disparar:", font=FONT_BOLD).pack(anchor="w", pady=(2, 4))

        self.scrollable_frame = ctk.CTkScrollableFrame(box_msgs, corner_radius=10)
        self.scrollable_frame.pack(fill="both", expand=True, pady=(0, 4))

        ctk.CTkButton(box_msgs, text="+ Criar Nova Mensagem", font=FONT_BOLD, corner_radius=8, height=34, command=self.add_snippet).pack(fill="x", pady=(0, 4))

        self.lbl_status = ctk.CTkLabel(self, text="● Pronto para disparo em lote (F8 Oculta/Mostra)", font=FONT_SMALL, text_color=("gray30", "gray70"))
        self.lbl_status.pack(anchor="w", padx=16, pady=(0, 10))

        self.registered_hotkeys = []
        self.apply_theme()
        self.update_categories_combo()
        self.update_parsed_numbers()

        if KEYBOARD_AVAILABLE:
            try:
                keyboard.add_hotkey('f8', self.toggle_visibility_safe)
                self.registered_hotkeys.append('f8')
            except: pass

        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def clear_recipients(self):
        self.txt_whatsapp_batch.delete("1.0", "end")
        self.update_parsed_numbers()

    def add_number_to_list(self):
        raw_input = self.txt_whatsapp_batch.get("1.0", "end-1c").strip()
        if not raw_input:
            messagebox.showwarning("Aviso", "Digite ou cole um número na caixa acima antes de importar.", parent=self)
            return

        new_numbers = []
        lines = raw_input.split('\n')
        for line in lines:
            clean_num = re.sub(r'\D', '', line)
            if clean_num:
                if len(clean_num) in (10, 11) and not clean_num.startswith('55'):
                    clean_num = '55' + clean_num
                if len(clean_num) >= 10 and clean_num not in new_numbers:
                    new_numbers.append(clean_num)

        if not new_numbers:
            messagebox.showwarning("Aviso", "Nenhum número válido encontrado no texto.", parent=self)
            return

        current_valid = self.extract_valid_numbers_from_box()
        combined = list(dict.fromkeys(current_valid + new_numbers))

        self.txt_whatsapp_batch.delete("1.0", "end")
        self.txt_whatsapp_batch.insert("1.0", "\n".join(combined))
        
        self.update_parsed_numbers()
        messagebox.showinfo("Sucesso", f"{len(new_numbers)} número(s) importado(s) para a lista de disparos!", parent=self)

    def extract_valid_numbers_from_box(self):
        raw_input = self.txt_whatsapp_batch.get("1.0", "end-1c").strip()
        if not raw_input:
            return []

        lines = raw_input.split('\n')
        valid_numbers = []

        for line in lines:
            clean_num = re.sub(r'\D', '', line)
            if clean_num:
                if len(clean_num) in (10, 11) and not clean_num.startswith('55'):
                    clean_num = '55' + clean_num
                if len(clean_num) >= 10 and clean_num not in valid_numbers:
                    valid_numbers.append(clean_num)
        return valid_numbers

    def import_spreadsheet(self):
        file_path = filedialog.askopenfilename(
            title="Selecionar Planilha de Contatos",
            filetypes=[
                ("Arquivos Excel e Texto", "*.xlsx;*.xls;*.csv;*.txt"),
                ("Excel Files", "*.xlsx;*.xls"),
                ("Text/CSV Files", "*.csv;*.txt"),
                ("Todos os arquivos", "*.*")
            ]
        )
        if not file_path:
            return

        extracted_texts = []
        try:
            if file_path.endswith(('.xlsx', '.xls')):
                if not EXCEL_AVAILABLE:
                    messagebox.showerror("Erro", "A biblioteca 'openpyxl' não está instalada.", parent=self)
                    return
                
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheet = wb.active
                
                target_col_idx = None
                for col_idx in range(1, sheet.max_column + 1):
                    cell_val = str(sheet.cell(row=1, column=col_idx).value or "").strip().lower()
                    if "telefone" in cell_val or "celular" in cell_val or "whats" in cell_val or "contato" in cell_val:
                        target_col_idx = col_idx
                        break
                
                if not target_col_idx and sheet.max_column > 0:
                    target_col_idx = 1
                
                if target_col_idx:
                    for row_idx in range(2, sheet.max_row + 1):
                        val = sheet.cell(row=row_idx, column=target_col_idx).value
                        if val is not None:
                            extracted_texts.append(str(val))
            else:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    lines = f.readlines()
                    if not lines:
                        return
                    for line in lines[1:]:
                        extracted_texts.append(line.strip())

            current_valid = self.extract_valid_numbers_from_box()
            
            new_file_numbers = []
            for text in extracted_texts:
                clean_num = re.sub(r'\D', '', text)
                if clean_num:
                    if len(clean_num) in (10, 11) and not clean_num.startswith('55'):
                        clean_num = '55' + clean_num
                    if len(clean_num) >= 10 and clean_num not in new_file_numbers:
                        new_file_numbers.append(clean_num)

            combined = list(dict.fromkeys(current_valid + new_file_numbers))
            
            self.txt_whatsapp_batch.delete("1.0", "end")
            self.txt_whatsapp_batch.insert("1.0", "\n".join(combined))
            self.update_parsed_numbers()
            
            messagebox.showinfo("Sucesso", f"Planilha importada!\n{len(new_file_numbers)} números adicionados à lista.", parent=self)

        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível ler o arquivo:\n{e}", parent=self)

    def extract_valid_numbers(self):
        return self.extract_valid_numbers_from_box()

    def update_parsed_numbers(self, event=None):
        for widget in self.scroll_column_view.winfo_children():
            widget.destroy()

        valid_numbers = self.extract_valid_numbers()
        self.lbl_count_valid.configure(text=f"{len(valid_numbers)} números")

        if not valid_numbers:
            lbl = ctk.CTkLabel(self.scroll_column_view, text="Nenhum número detectado.", font=FONT_SMALL, text_color="gray")
            lbl.pack(anchor="w", padx=4)
        else:
            for idx, num in enumerate(valid_numbers, 1):
                row = ctk.CTkFrame(self.scroll_column_view, fg_color="transparent")
                row.pack(fill="x", pady=1)
                ctk.CTkLabel(row, text=f"{idx}.", font=FONT_BADGE, width=25, text_color="gray").pack(side="left")
                ctk.CTkLabel(row, text=f"+{num}", font=FONT_MAIN).pack(side="left")

    def toggle_visibility_safe(self):
        self.after(0, self.toggle_visibility)

    def toggle_visibility(self):
        if self.is_visible:
            self.withdraw()
            self.is_visible = False
        else:
            self.deiconify()
            self.attributes('-topmost', True)
            self.is_visible = True

    def show_history(self):
        if self.active_history_dialog is None or not self.active_history_dialog.winfo_exists():
            HistoryDialog(self)
        else:
            self.active_history_dialog.focus()

    def apply_theme(self):
        if self.current_theme == "light": ctk.set_appearance_mode("Light")
        else: ctk.set_appearance_mode("Dark")

    def toggle_theme(self):
        self.current_theme = "light" if self.current_theme == "dark" else "dark"
        self.data["settings"]["theme"] = self.current_theme
        save_data(self.data)
        self.apply_theme()
        idx = self.get_current_cat_index()
        if idx >= 0: self.load_snippets(idx)

    def update_categories_combo(self, select_index=0):
        categories = self.data.get("categories", [])
        cat_names = [c["name"] for c in categories]
        if cat_names:
            self.combo_cat.configure(values=cat_names)
            select_index = min(select_index, len(cat_names) - 1)
            self.combo_cat.set(cat_names[select_index])
            self.load_snippets(select_index)
        else:
            self.combo_cat.configure(values=["Sem categorias"])
            self.combo_cat.set("Sem categorias")
            self.clear_snippets_view()

    def get_current_cat_index(self):
        cat_name = self.combo_cat.get()
        for idx, cat in enumerate(self.data.get("categories", [])):
            if cat["name"] == cat_name: return idx
        return -1

    def on_category_change(self, choice):
        idx = self.get_current_cat_index()
        if idx >= 0: self.load_snippets(idx)

    def clear_snippets_view(self):
        if KEYBOARD_AVAILABLE:
            for sc in self.registered_hotkeys:
                if sc != 'f8':
                    try: keyboard.remove_hotkey(sc)
                    except: pass
            self.registered_hotkeys = [sc for sc in self.registered_hotkeys if sc == 'f8']
        for widget in self.scrollable_frame.winfo_children(): widget.destroy()

    def load_snippets(self, cat_index):
        self.clear_snippets_view()
        categories = self.data.get("categories", [])
        if not categories or cat_index < 0 or cat_index >= len(categories): return
        snippets = categories[cat_index].get("snippets", [])

        for idx, item in enumerate(snippets):
            shortcut = item.get("shortcut", "").upper()
            card_frame = ctk.CTkFrame(self.scrollable_frame, corner_radius=10, fg_color=("gray85", "gray20"))
            card_frame.pack(fill="x", pady=4, expand=True)

            btn_main = ctk.CTkButton(
                card_frame, text=f"🚀 {item['label']}", font=FONT_BOLD, fg_color="transparent", 
                text_color=("gray10", "gray90"), anchor="w", height=38,
                command=lambda b=item["body"], l=item["label"]: self.process_snippet(b, l)
            )
            btn_main.pack(side="left", fill="x", expand=True, padx=(6, 0))

            if shortcut:
                lbl_sc = ctk.CTkLabel(card_frame, text=shortcut, font=FONT_BADGE, fg_color=("gray75", "gray30"), text_color=("gray10", "gray90"), corner_radius=6, width=60, height=22)
                lbl_sc.pack(side="left", padx=(0, 6))

            btn_edit = ctk.CTkButton(card_frame, text="✎", font=("Segoe UI", 12, "bold"), fg_color="transparent", text_color=("gray20", "gray80"), width=30, height=30, command=lambda i=idx: self.edit_snippet(i))
            btn_edit.pack(side="left")

            btn_del = ctk.CTkButton(card_frame, text="✕", font=("Segoe UI", 11, "bold"), fg_color="transparent", text_color="#e74c3c", width=30, height=30, command=lambda i=idx: self.delete_snippet(i))
            btn_del.pack(side="left", padx=(0, 4))

            if shortcut and KEYBOARD_AVAILABLE:
                try:
                    sc_lower = shortcut.lower()
                    keyboard.add_hotkey(sc_lower, lambda b=item["body"], l=item["label"]: self.after(0, lambda: self.process_snippet(b, l)))
                    if sc_lower not in self.registered_hotkeys:
                        self.registered_hotkeys.append(sc_lower)
                except: pass

    def add_category(self):
        new_name = simpledialog.askstring("Novo Grupo", "Nome do grupo de mensagens:", parent=self)
        if new_name and new_name.strip():
            cat_id = re.sub(r'[^a-zA-Z0-9]', '-', new_name.lower())
            self.data.setdefault("categories", []).append({"id": cat_id, "name": new_name.strip(), "snippets": []})
            save_data(self.data)
            self.update_categories_combo(len(self.data["categories"]) - 1)

    def add_snippet(self):
        cat_idx = self.get_current_cat_index()
        if cat_idx < 0: return messagebox.showwarning("Aviso", "Crie ou selecione um grupo primeiro.", parent=self)
        dialog = EditSnippetDialog(self)
        self.wait_window(dialog)
        if dialog.saved:
            snip_id = re.sub(r'[^a-zA-Z0-9]', '-', dialog.label_val.lower())
            self.data["categories"][cat_idx]["snippets"].append({"id": snip_id, "label": dialog.label_val, "shortcut": dialog.shortcut_val, "body": dialog.body_val})
            save_data(self.data)
            self.load_snippets(cat_idx)

    def edit_snippet(self, snip_idx):
        cat_idx = self.get_current_cat_index()
        snip = self.data["categories"][cat_idx]["snippets"][snip_idx]
        dialog = EditSnippetDialog(self, label=snip["label"], body=snip["body"], shortcut=snip.get("shortcut", ""))
        self.wait_window(dialog)
        if dialog.saved:
            snip["label"] = dialog.label_val
            snip["shortcut"] = dialog.shortcut_val
            snip["body"] = dialog.body_val
            save_data(self.data)
            self.load_snippets(cat_idx)

    def delete_snippet(self, snip_idx):
        cat_idx = self.get_current_cat_index()
        snip_label = self.data["categories"][cat_idx]["snippets"][snip_idx]["label"]
        if messagebox.askyesno("Excluir", f"Excluir a mensagem '{snip_label}'?", parent=self):
            del self.data["categories"][cat_idx]["snippets"][snip_idx]
            save_data(self.data)
            self.load_snippets(cat_idx)

    def log_action(self, label, phone=""):
        today = datetime.now().strftime("%Y-%m-%d")
        history_data = load_history()
        
        if today not in history_data:
            history_data[today] = {"total": 0, "snippets": {}, "numbers": []}
            
        history_data[today]["total"] += 1
        history_data[today]["snippets"][label] = history_data[today]["snippets"].get(label, 0) + 1
        
        if phone:
            time_str = datetime.now().strftime("%H:%M")
            history_data[today]["numbers"].append(f"[{time_str}] {phone} - {label}")
            
        save_history(history_data)
        if self.active_history_dialog and self.active_history_dialog.winfo_exists():
            self.active_history_dialog.refresh_data()

    def process_snippet(self, body, label):
        now = datetime.now()
        if now.hour < 12: saudacao = "Bom dia"
        elif now.hour < 18: saudacao = "Boa tarde"
        else: saudacao = "Boa noite"

        body = body.replace("[HOJE]", now.strftime("%d/%m/%Y"))
        body = body.replace("[AMANHA]", (now + timedelta(days=1)).strftime("%d/%m/%Y"))
        body = body.replace("[SAUDACAO]", saudacao)

        has_placeholders = bool(re.search(r'\[(.*?)\]', body))
        if has_placeholders:
            dialog = FillSnippetDialog(self, label, body)
            self.wait_window(dialog)
            if dialog.result_text:
                self.execute_batch_dispatch(dialog.result_text, label)
        else:
            self.execute_batch_dispatch(body, label)

    def execute_batch_dispatch(self, text, label):
        valid_numbers = self.extract_valid_numbers()

        if not valid_numbers:
            messagebox.showwarning("Atenção", "Insira pelo menos um número válido na lista de destinatários.", parent=self)
            return

        encoded_text = urllib.parse.quote(text)
        
        for i, phone in enumerate(valid_numbers):
            url = f"whatsapp://send?phone={phone}&text={encoded_text}"
            webbrowser.open(url)
            self.log_action(label, phone)
            
            if i < len(valid_numbers) - 1:
                time.sleep(1.5)

        self.lbl_status.configure(text=f"⚡ Sucesso: {len(valid_numbers)} mensagens disparadas!", text_color="#2ecc71")
        self.after(4000, lambda: self.lbl_status.configure(text="● Pronto para disparo em lote (F8 Oculta/Mostra)", text_color=("gray30", "gray70")))

    def on_close(self):
        if KEYBOARD_AVAILABLE:
            for sc in self.registered_hotkeys:
                try: keyboard.remove_hotkey(sc)
                except: pass
        self.destroy()

if __name__ == "__main__":
    app = MassSenderApp()
    app.mainloop()